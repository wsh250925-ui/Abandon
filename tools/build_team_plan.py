"""Build the editable team plan from docs/TEAM_PLAN.md using python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / '急救学习小程序_三人分工与开发安排_V0.1.docx'

def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
font(doc.styles['Normal'], 11)
doc.styles['Normal'].paragraph_format.line_spacing = Pt(15)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
for name, size in [('Title', 22), ('Heading 1', 16), ('Heading 2', 12)]:
    font(doc.styles[name], size, True)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
for style in doc.styles:
    props = style.element.find(qn('w:pPr'))
    if props is not None:
        for el in list(props):
            if el.tag == qn('w:pBdr'): props.remove(el)
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); props.append(snap)
for grid in list(section._sectPr.findall(qn('w:docGrid'))): section._sectPr.remove(grid)
doc.core_properties.title = '健康学院急救学习小程序三人分工与开发安排'
doc.core_properties.author = '急救学习项目组'
doc.core_properties.subject = '三人成员责任 任务拆分 开发排期与验收'

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('第 ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
footer.add_run(' 页').font.size = Pt(9)

def table(rows):
    count = len(rows[0]); t = doc.add_table(rows=1, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if rows[0][0] == '编号': fractions = [.08, .50, .24, .18]
    elif count == 5: fractions = [.12, .16, .17, .18, .37]
    elif count == 4: fractions = [.43, .12, .22, .23]
    elif count == 3: fractions = [.24, .39, .37]
    else: fractions = [1/count]*count
    if rows[0][0] == '阶段': fractions = [.14,.19,.19,.22,.26]
    for col, f in zip(t.columns, fractions): col.width = Inches(7.1*f)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement('w:cantSplit'); t.rows[i]._tr.get_or_add_trPr().append(no_split)
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Inches(7.1*fractions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '1D4ED8' if i == 0 else ('F3F6FB' if i % 2 == 0 else 'FFFFFF')); props.append(shade)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, value_margin in [('top',65),('bottom',65),('left',85),('right',85)]:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),str(value_margin)); el.set(qn('w:type'),'dxa'); margins.append(el)
            props.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.line_spacing=Pt(14); p.paragraph_format.space_after=Pt(0)
            if count == 5 and rows[0][0] == '成员': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif count == 4 and rows[0][0] == '工作项' and j == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(value); r.font.size=Pt(10.5); r.bold=(i==0)
            if i==0: r.font.color.rgb=RGBColor(255,255,255)
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing=Pt(4)

lines = (ROOT/'docs'/'TEAM_PLAN.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if line == '<!-- page -->': doc.add_page_break()
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            parts=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',x) for x in parts): rows.append(parts)
            i+=1
        table(rows); continue
    elif line.startswith('# '): doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '): doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '): doc.add_paragraph(line[4:], 'Heading 2')
    elif line:
        p=doc.add_paragraph(line)
        if line.startswith('初版方案'): p.runs[0].font.size=Pt(10)
    i+=1
doc.save(OUT)
print(OUT)
